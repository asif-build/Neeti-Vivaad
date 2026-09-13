import json
from unittest.mock import patch

from django.test import TestCase
from rest_framework.test import APIClient

from core.models import User
from assessment.generator import QuizGenerationError, _validate_questions


SOURCE_TEXT = (
    "The policy requires encrypted storage on every field device. "
    "A verified backup must be completed at the end of each working day. "
    "Supervisors must review failed backups before new records are collected. "
    "These controls apply to every district office and all mobile field teams."
)


def generated_payload():
    return {
        "questions": [
            {
                "question": "What storage control does the policy require for field devices?",
                "source_citation": "The policy requires encrypted storage on every field device.",
                "explanation": "The source directly states that every field device requires encrypted storage.",
                "options": [
                    {"text": "Encrypted storage", "is_correct": True},
                    {"text": "Public cloud storage", "is_correct": False},
                    {"text": "Removable storage only", "is_correct": False},
                    {"text": "No device storage", "is_correct": False},
                ],
            },
            {
                "question": "When must the verified backup be completed under the policy?",
                "source_citation": "A verified backup must be completed at the end of each working day.",
                "explanation": "The cited sentence specifies the end of each working day as the deadline.",
                "options": [
                    {"text": "At the end of each working day", "is_correct": True},
                    {"text": "At the end of each month", "is_correct": False},
                    {"text": "Before each working day", "is_correct": False},
                    {"text": "Only after an incident", "is_correct": False},
                ],
            },
        ]
    }


class QuizValidationTests(TestCase):
    def test_rejects_citation_not_found_in_source(self):
        payload = generated_payload()
        payload["questions"][0]["source_citation"] = "This sentence is not in the source."

        with self.assertRaises(QuizGenerationError):
            _validate_questions(payload, SOURCE_TEXT, 2)


class QuizApiTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="official", password="secret")
        self.client = APIClient()
        self.client.force_authenticate(self.user)

    @patch("assessment.generator.generate_text")
    def test_upload_generate_and_submit_uses_generated_questions(self, generate_text):
        generate_text.return_value = json.dumps(generated_payload())

        upload = self.client.post(
            "/api/assessment/upload/",
            {"title": "Field controls", "text": SOURCE_TEXT},
            format="json",
        )
        self.assertEqual(upload.status_code, 201)

        generation = self.client.post(
            "/api/assessment/generate-quiz/",
            {"document_id": upload.data["document_id"], "num_questions": 2},
            format="json",
        )
        self.assertEqual(generation.status_code, 201)
        self.assertEqual(len(generation.data["questions"]), 2)
        self.assertEqual(
            generation.data["questions"][0]["question"],
            generated_payload()["questions"][0]["question"],
        )

        answers = {
            str(question["id"]): question["options"][0]["id"]
            for question in generation.data["questions"]
        }
        submission = self.client.post(
            "/api/assessment/submit-quiz/",
            {"quiz_id": generation.data["quiz_id"], "answers": answers},
            format="json",
        )
        self.assertEqual(submission.status_code, 200)
        self.assertEqual(submission.data["score_percentage"], 100.0)

    def test_upload_rejects_empty_content(self):
        response = self.client.post(
            "/api/assessment/upload/",
            {"title": "Empty", "text": ""},
            format="json",
        )
        self.assertEqual(response.status_code, 400)


class DocumentExtractionTests(TestCase):
    def test_docx_in_memory_extraction(self):
        import io
        import docx
        from assessment.extraction import process_document_source

        doc = docx.Document()
        doc.add_heading("National Statistical Protocol", level=1)
        doc.add_paragraph("All microdata must adhere to k-anonymity where k is at least 5 before dissemination.")
        doc.add_paragraph("Supervisory audit must occur within 48 hours of field collection completion.")
        buf = io.BytesIO()
        doc.save(buf)
        file_bytes = buf.getvalue()

        result = process_document_source(file_bytes, "protocol.docx")
        self.assertEqual(result["file_type"], "DOCX")
        self.assertTrue(len(result["chunks"]) >= 1)
        self.assertIn("k-anonymity", result["extracted_text"])

    def test_rejects_unsupported_file_extension(self):
        from assessment.extraction import process_document_source, DocumentExtractionError
        with self.assertRaises(DocumentExtractionError):
            process_document_source(b"binary content", "script.exe")

    def test_rejects_unreadable_or_empty_text(self):
        from assessment.extraction import process_document_source, DocumentExtractionError
        with self.assertRaises(DocumentExtractionError):
            process_document_source(b"too short", "note.txt")


class StudioLifecycleAndSecurityTests(TestCase):
    def setUp(self):
        self.author = User.objects.create_user(username="author1", email="author1@gov.in", password="password")
        self.intruder = User.objects.create_user(username="intruder", email="intruder@gov.in", password="password")
        self.client_author = APIClient()
        self.client_author.force_authenticate(self.author)
        self.client_intruder = APIClient()
        self.client_intruder.force_authenticate(self.intruder)

    def test_studio_workflow_and_idor_protection(self):
        policy_text = (
            "Section 1. Microdata Security Standards:\n"
            "Field enumerators must retain encrypted records on secure tablets.\n"
            "All national sample statistical collections must maintain a minimum confidence interval of 95 percent.\n"
            "Automated anomaly detection must flag duplicate household records within 24 hours of submission.\n"
            "Data dissemination must undergo k-anonymity and differential privacy noise addition."
        )

        # 1. Author uploads document
        up_res = self.client_author.post("/api/assessment/documents/", {
            "title": "MoSPI Microdata Standards",
            "text": policy_text
        }, format="json")
        self.assertEqual(up_res.status_code, 201)
        doc_id = up_res.data["document_id"]

        # 2. Intruder attempts IDOR to generate quiz from Author's document
        intruder_gen = self.client_intruder.post("/api/assessment/studio/generate/", {
            "document_id": doc_id,
            "num_questions": 2
        }, format="json")
        self.assertEqual(intruder_gen.status_code, 404)

        # 3. Author generates draft quiz
        gen_res = self.client_author.post("/api/assessment/studio/generate/", {
            "document_id": doc_id,
            "title": "MoSPI Security Knowledge Check",
            "num_questions": 2,
            "difficulty": "Intermediate"
        }, format="json")
        self.assertEqual(gen_res.status_code, 201)
        quiz_id = gen_res.data["quiz_id"]
        self.assertEqual(gen_res.data["status"], "DRAFT")
        self.assertEqual(len(gen_res.data["questions"]), 2)

        # 4. Draft quiz is NOT discoverable in public/learner catalog
        cat_res = self.client_author.get("/api/assessment/checks/")
        self.assertEqual(cat_res.status_code, 200)
        self.assertFalse(any(c["id"] == quiz_id for c in cat_res.data["checks"]))

        # 5. Intruder attempts to edit author's quiz (IDOR)
        intruder_edit = self.client_intruder.patch(f"/api/assessment/studio/{quiz_id}/", {
            "title": "Hacked Title"
        }, format="json")
        self.assertEqual(intruder_edit.status_code, 403)

        # 6. Author adds a manual question
        add_res = self.client_author.post(f"/api/assessment/studio/{quiz_id}/questions/", {
            "action": "add",
            "question_text": "What is the mandatory threshold for household record duplicate checks?",
            "question_type": "MCQ",
            "source_page": 1,
            "source_section": "Section 1",
            "evidence_text": "Automated anomaly detection must flag duplicate household records within 24 hours.",
            "explanation": "Official guidelines mandate automated flagging within 24 hours.",
            "options": [
                {"text": "Within 24 hours", "is_correct": True},
                {"text": "Within 30 days", "is_correct": False},
                {"text": "At annual audit", "is_correct": False}
            ]
        }, format="json")
        self.assertEqual(add_res.status_code, 201)

        # 7. Author publishes the Knowledge Check
        pub_res = self.client_author.post(f"/api/assessment/studio/{quiz_id}/publish/")
        self.assertEqual(pub_res.status_code, 200)
        self.assertEqual(pub_res.data["status"], "PUBLISHED")

        # 8. Now published check is visible in catalog
        cat_pub_res = self.client_author.get("/api/assessment/checks/")
        self.assertTrue(any(c["id"] == quiz_id for c in cat_pub_res.data["checks"]))

        # 9. Learner takes published check and receives source-backed feedback
        detail_res = self.client_intruder.get(f"/api/assessment/checks/{quiz_id}/")
        self.assertEqual(detail_res.status_code, 200)
        q_list = detail_res.data["questions"]
        self.assertTrue(len(q_list) >= 3)

        # Build answers: pick first option for each
        answers = {str(q["id"]): q["options"][0]["id"] for q in q_list}
        sub_res = self.client_intruder.post(f"/api/assessment/checks/{quiz_id}/submit/", {
            "answers": answers
        }, format="json")
        self.assertEqual(sub_res.status_code, 200)
        self.assertIn("score_percentage", sub_res.data)
        self.assertIn("detailed_results", sub_res.data)
        self.assertIn("what_you_did_well", sub_res.data)
        self.assertIn("keep_practising", sub_res.data)
        # Check source page in result
        self.assertEqual(sub_res.data["detailed_results"][0]["source_page"], 1)
